#!/usr/bin/env python3
"""Create the empty local metadata workbook expected by the histogram script."""
import argparse
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def main():
    default = (Path(__file__).resolve().parent.parent / "scripts" /
               "FinalPythonScripts" / "subject_info_NIH.xlsx")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=default)
    args = parser.parse_args()
    target = args.output.resolve()
    if target.exists():
        parser.error(f"Refusing to overwrite existing file: {target}")
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Subject_Info"
    sheet.append(["Subject_Number", "Group", "Subgroup", "Threshold_Score"])
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244660")
    for column, width in zip("ABCD", [23, 20, 20, 23]):
        sheet.column_dimensions[column].width = width
    # Exclusive creation also prevents overwriting a file created after the check.
    with target.open("xb") as output:
        workbook.save(output)
    workbook.close()
    print(f"Created empty subject information template: {target}")


if __name__ == "__main__":
    main()
